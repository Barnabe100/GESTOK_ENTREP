"""Export Excel (.xlsx) générique, réutilisable par tous les rapports.

Même rôle que ``csv_export.py``/``pdf_export.py`` : fonction utilitaire pure,
sans vérification de permission (déjà effectuée par l'appelant au moment
d'obtenir les données du rapport) et sans connaissance métier. Utilise
``openpyxl`` (seule nouvelle dépendance ajoutée par ce lot — voir
requirements.txt).

Contrairement à ``export_rows_to_csv``, dont les lignes sont déjà formatées
en texte par l'appelant (convention CSV : tout est du texte de toute façon),
``export_rows_to_xlsx`` attend des valeurs réellement typées (texte,
``int``, ``float`` ou ``Decimal``) afin que les colonnes numériques
(montants, quantités) restent triables et calculables dans Excel — un
fichier CSV n'offrant aucun avantage à le faire, ce n'est qu'avec ce module
que cette distinction existe.
"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Optional, Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.utils.exceptions import ValidationError

CellValue = Union[str, int, float, Decimal, None]

_MIN_COLUMN_WIDTH = 10
_MAX_COLUMN_WIDTH = 60


def export_rows_to_xlsx(
    file_path: Union[str, Path],
    headers: Sequence[str],
    rows: Sequence[Sequence[CellValue]],
    column_formats: Optional[Sequence[Optional[str]]] = None,
) -> None:
    """Écrit ``headers``/``rows`` dans un classeur Excel à ``file_path``.

    ``column_formats`` donne, pour chaque colonne (même ordre que
    ``headers``), un format de cellule Excel optionnel (ex. ``"#,##0.00"``
    pour un montant) ; ``None`` à un index laisse Excel appliquer son format
    général à cette colonne. Lève :class:`ValidationError` si le fichier ne
    peut pas être écrit (dossier inaccessible, ...) ou si l'export n'a
    produit aucun fichier exploitable — même contrat que
    ``export_document_to_pdf``."""
    file_path = Path(file_path)
    try:
        file_path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ValidationError(f"Impossible de créer le dossier de destination : {exc}") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rapport"

    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    rows = list(rows)
    for row in rows:
        sheet.append(list(row))

    if column_formats:
        for col_index, number_format in enumerate(column_formats, start=1):
            if not number_format:
                continue
            column_letter = get_column_letter(col_index)
            for row_index in range(2, len(rows) + 2):
                sheet[f"{column_letter}{row_index}"].number_format = number_format

    for col_index, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows:
            if col_index - 1 < len(row) and row[col_index - 1] is not None:
                longest = max(longest, len(str(row[col_index - 1])))
        sheet.column_dimensions[get_column_letter(col_index)].width = max(
            _MIN_COLUMN_WIDTH, min(_MAX_COLUMN_WIDTH, longest + 2)
        )

    try:
        workbook.save(file_path)
    except OSError as exc:
        raise ValidationError(f"Impossible d'écrire le fichier Excel : {exc}") from exc

    if not file_path.exists() or file_path.stat().st_size == 0:
        raise ValidationError(f"L'export Excel a échoué : le fichier « {file_path} » n'a pas été créé.")
