from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.utils.exceptions import ValidationError
from app.utils.xlsx_export import export_rows_to_xlsx


def test_export_writes_valid_xlsx_file(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(out, ["Référence", "Stock"], [["ART-1", 10], ["ART-2", 5]])

    assert out.exists()
    assert out.stat().st_size > 0
    assert out.read_bytes()[:2] == b"PK"  # signature ZIP (format .xlsx)


def test_export_writes_correct_headers_and_data(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(out, ["Référence", "Stock"], [["ART-1", 10], ["ART-2", 5]])

    workbook = openpyxl.load_workbook(out)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Référence", "Stock"]
    assert [cell.value for cell in sheet[2]] == ["ART-1", 10]
    assert [cell.value for cell in sheet[3]] == ["ART-2", 5]


def test_header_row_is_bold(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(out, ["Référence"], [["ART-1"]])

    workbook = openpyxl.load_workbook(out)
    sheet = workbook.active
    assert sheet["A1"].font.bold is True
    assert sheet["A2"].font.bold is not True


def test_numeric_values_are_stored_as_real_numbers_not_text(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(out, ["Désignation", "Montant", "Quantité"], [["Article", Decimal("1234.50"), Decimal("10")]])

    workbook = openpyxl.load_workbook(out)
    sheet = workbook.active
    assert sheet["A2"].value == "Article"
    assert sheet["A2"].data_type == "s"
    assert float(sheet["B2"].value) == 1234.50
    assert sheet["B2"].data_type == "n"
    assert float(sheet["C2"].value) == 10
    assert sheet["C2"].data_type == "n"


def test_column_formats_apply_number_format_to_data_cells_only(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(
        out,
        ["Désignation", "Montant"],
        [["A", Decimal("1500")], ["B", Decimal("2500")]],
        column_formats=[None, "#,##0.00"],
    )

    workbook = openpyxl.load_workbook(out)
    sheet = workbook.active
    assert sheet["B2"].number_format == "#,##0.00"
    assert sheet["B3"].number_format == "#,##0.00"
    assert sheet["A2"].number_format != "#,##0.00"
    # La ligne d'en-tête ne porte jamais le format numérique de la colonne.
    assert sheet["B1"].number_format != "#,##0.00"


def test_export_with_no_rows_still_produces_valid_file_with_only_headers(tmp_path) -> None:
    out = tmp_path / "out.xlsx"
    export_rows_to_xlsx(out, ["Référence", "Stock"], [])

    workbook = openpyxl.load_workbook(out)
    sheet = workbook.active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == ["Référence", "Stock"]


def test_export_creates_missing_parent_directory(tmp_path) -> None:
    out = tmp_path / "not_yet_created" / "out.xlsx"
    export_rows_to_xlsx(out, ["A"], [["1"]])

    assert out.exists()


def test_export_rejects_destination_whose_parent_is_a_file(tmp_path) -> None:
    blocked = tmp_path / "obstacle"
    blocked.write_text("fichier, pas dossier")
    out = blocked / "out.xlsx"

    with pytest.raises(ValidationError):
        export_rows_to_xlsx(out, ["A"], [["1"]])


def test_export_accepts_string_path(tmp_path) -> None:
    out = str(tmp_path / "out.xlsx")
    export_rows_to_xlsx(out, ["A"], [["1"]])

    assert Path(out).exists()
