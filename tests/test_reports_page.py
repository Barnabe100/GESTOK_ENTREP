from datetime import date
from decimal import Decimal

import pytest
from PySide6.QtWidgets import QMessageBox

from app.services.entries.entry_service import EntreeLigneInput
from app.services.sales.sale_service import VenteLigneInput
from app.views.optional_date_edit import OptionalDateEdit
from app.views.pages.reports_page import ReportsPage


@pytest.fixture(autouse=True)
def _no_blocking_dialogs(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.views.pages.reports_page.QMessageBox.information", lambda *a, **k: None)
    monkeypatch.setattr("app.views.pages.reports_page.QMessageBox.warning", lambda *a, **k: None)


def _build_page(stack) -> ReportsPage:
    return ReportsPage(
        stack.reports, stack.categories, stack.articles, stack.exit_reasons, stack.permissions
    )


def _make_article(stack, reference="ART-1", stock_initial=Decimal("50"), stock_min=Decimal("0")):
    category = stack.categories.create_category(f"Cat-{reference}")
    return stack.articles.create_article(
        reference, "Article de test", category.id, "unité",
        Decimal("100"), Decimal("150"), stock_min, stock_initial=stock_initial,
    )


def test_default_report_is_stock_state(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack)

    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.report_combo.currentText() == "État du stock"
    assert page.table.rowCount() == 1


def test_stock_state_shows_valeur_stock_column(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    headers = [page.table.horizontalHeaderItem(i).text() for i in range(page.table.columnCount())]
    assert "Valeur stock" in headers


def test_stock_state_formats_quantities_without_unnecessary_decimals(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"), stock_min=Decimal("2"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    headers = [page.table.horizontalHeaderItem(i).text() for i in range(page.table.columnCount())]
    stock_col = headers.index("Stock actuel")
    stock_min_col = headers.index("Stock min")
    assert page.table.item(0, stock_col).text() == "10"  # jamais "10.000"
    assert page.table.item(0, stock_min_col).text() == "2"


def test_stock_state_preserves_real_decimal_stock(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10.5"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    headers = [page.table.horizontalHeaderItem(i).text() for i in range(page.table.columnCount())]
    stock_col = headers.index("Stock actuel")
    assert page.table.item(0, stock_col).text() == "10,5"


def test_switching_to_low_stock_report_updates_table(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, reference="ART-LOW", stock_initial=Decimal("1"), stock_min=Decimal("10"))
    _make_article(stack, reference="ART-OK", stock_initial=Decimal("50"), stock_min=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Stock faible")

    references = {page.table.item(row, 0).text() for row in range(page.table.rowCount())}
    assert references == {"ART-LOW"}


def test_out_of_stock_report(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, reference="ART-ZERO", stock_initial=Decimal("0"))
    _make_article(stack, reference="ART-POS", stock_initial=Decimal("5"))

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Ruptures")

    references = {page.table.item(row, 0).text() for row in range(page.table.rowCount())}
    assert references == {"ART-ZERO"}


def test_valorisation_report_shows_total_in_summary(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Valorisation")

    assert "Valeur totale" in page.summary_label.text()


def test_movements_report_lists_entries(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    supplier = stack.suppliers.create_supplier("F")
    article = _make_article(stack, stock_initial=Decimal("0"))
    entry = stack.entries.create_entry(supplier.id, date(2026, 1, 1), [EntreeLigneInput(article.id, Decimal("10"), Decimal("100"))])
    stack.entries.validate_entry(entry.id)

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Mouvements")

    assert page.table.rowCount() == 1


def test_movements_report_filters_by_article(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    category = stack.categories.create_category("Boissons")
    article_1 = stack.articles.create_article(
        "ART-1", "A", category.id, "unité", Decimal("100"), Decimal("150"), Decimal("0"), stock_initial=Decimal("10")
    )
    article_2 = stack.articles.create_article(
        "ART-2", "B", category.id, "unité", Decimal("100"), Decimal("150"), Decimal("0"), stock_initial=Decimal("20")
    )

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Mouvements")

    index = page.article_combo.findData(article_1.id)
    page.article_combo.setCurrentIndex(index)
    page.refresh()

    assert page.table.rowCount() == 1
    assert page.table.item(0, 1).text() == "ART-1"


def test_entries_report_lists_and_filters_by_statut(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    supplier = stack.suppliers.create_supplier("F")
    article = _make_article(stack)
    stack.entries.create_entry(supplier.id, date(2026, 1, 1), [])
    to_validate = stack.entries.create_entry(supplier.id, date(2026, 1, 1), [EntreeLigneInput(article.id, Decimal("5"), Decimal("100"))])
    stack.entries.validate_entry(to_validate.id)

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Entrées")

    assert page.table.rowCount() == 2

    index = page.statut_combo.findText("Validée")
    page.statut_combo.setCurrentIndex(index)
    page.refresh()

    assert page.table.rowCount() == 1
    assert page.table.item(0, 0).text() == to_validate.numero


def test_sales_report_shows_historized_total(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    article = _make_article(stack)
    stack.sales.create_sale(date(2026, 1, 1), [VenteLigneInput(article.id, Decimal("2"), Decimal("150"))])

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Ventes")

    assert page.table.rowCount() == 1
    assert "300" in page.table.item(0, 4).text()


def test_inventories_report_shows_ecart_counts(qtbot, login_as) -> None:
    from app.services.inventory.inventory_service import InventaireLigneInput

    stack, _ = login_as("Administrateur")
    article = _make_article(stack, stock_initial=Decimal("100"))
    stack.inventory.create_inventory(date(2026, 1, 1), [InventaireLigneInput(article.id, Decimal("97"))])

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Inventaires")

    assert page.table.rowCount() == 1
    assert page.table.item(0, 6).text() == "1"  # écarts négatifs


def test_reset_filters_clears_search_and_refreshes(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, reference="ART-EAU")
    _make_article(stack, reference="ART-RIZ")

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.search_edit.setText("EAU")
    page.refresh()
    assert page.table.rowCount() == 1

    page._on_reset_filters_clicked()

    assert page.search_edit.text() == ""
    assert page.table.rowCount() == 2


def test_export_button_disabled_without_permission(qtbot, login_as) -> None:
    stack, _ = login_as("Consultation")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.export_button.isEnabled() is False


def test_export_button_enabled_with_permission(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.export_button.isEnabled() is True


def test_export_to_writes_csv_file(qtbot, login_as, tmp_path) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    out_file = tmp_path / "export.csv"
    result = page._export_to(str(out_file))

    assert result is True
    assert out_file.exists()
    content = out_file.read_text(encoding="utf-8-sig")
    assert "Référence" in content


# -- export Excel -------------------------------------------------------------


def test_export_xlsx_button_disabled_without_permission(qtbot, login_as) -> None:
    stack, _ = login_as("Consultation")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.export_xlsx_button.isEnabled() is False


def test_export_xlsx_button_enabled_with_permission(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.export_xlsx_button.isEnabled() is True


def test_export_xlsx_to_writes_valid_file_without_affecting_csv(qtbot, login_as, tmp_path) -> None:
    import openpyxl

    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    out_file = tmp_path / "export.xlsx"
    result = page._export_xlsx_to(str(out_file))

    assert result is True
    assert out_file.exists()
    workbook = openpyxl.load_workbook(out_file)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert "Référence" in headers
    assert "Valeur stock" in headers

    # L'export CSV reste utilisable sans aucune interférence (deux exports
    # indépendants, voir Lot 6).
    out_csv = tmp_path / "export.csv"
    assert page._export_to(str(out_csv)) is True
    assert out_csv.exists()


def test_export_xlsx_stores_amounts_and_quantities_as_real_numbers(qtbot, login_as, tmp_path) -> None:
    import openpyxl

    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10.5"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    out_file = tmp_path / "export.xlsx"
    page._export_xlsx_to(str(out_file))

    workbook = openpyxl.load_workbook(out_file)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    stock_col = headers.index("Stock actuel") + 1
    cmup_col = headers.index("CMUP") + 1
    valeur_col = headers.index("Valeur stock") + 1

    assert float(sheet.cell(row=2, column=stock_col).value) == 10.5
    assert sheet.cell(row=2, column=stock_col).data_type == "n"
    assert sheet.cell(row=2, column=cmup_col).data_type == "n"
    assert sheet.cell(row=2, column=valeur_col).data_type == "n"
    # Montants : format de cellule Excel dédié (pas le format général) —
    # devise par défaut XOF, sans décimales (voir money_decimal_places).
    assert sheet.cell(row=2, column=cmup_col).number_format == "#,##0"
    assert sheet.cell(row=2, column=valeur_col).number_format == "#,##0"


def test_export_xlsx_headers_match_displayed_table(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    _make_article(stack, stock_initial=Decimal("10"))

    page = _build_page(stack)
    qtbot.addWidget(page)

    displayed_headers = [page.table.horizontalHeaderItem(i).text() for i in range(page.table.columnCount())]
    assert page._current_headers == displayed_headers
    assert len(page._current_rows_raw) == page.table.rowCount()
    assert len(page._current_column_formats) == len(displayed_headers)


def test_export_xlsx_with_empty_report_still_succeeds(qtbot, login_as, tmp_path) -> None:
    import openpyxl

    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Stock faible")  # aucun article créé -> rapport vide

    assert page.table.rowCount() == 0

    out_file = tmp_path / "export.xlsx"
    result = page._export_xlsx_to(str(out_file))

    assert result is True
    workbook = openpyxl.load_workbook(out_file)
    sheet = workbook.active
    assert sheet.max_row == 1  # en-têtes uniquement


def test_export_xlsx_with_no_data_shows_information_message_and_writes_nothing(qtbot, login_as, tmp_path) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)
    page._current_headers = []

    out_file = tmp_path / "export.xlsx"
    page._on_export_xlsx_clicked()

    assert not out_file.exists()


def test_export_xlsx_reports_errors_from_denied_permission(qtbot, login_as, tmp_path) -> None:
    stack, _ = login_as("Consultation")
    page = _build_page(stack)
    qtbot.addWidget(page)

    out_file = tmp_path / "export.xlsx"
    result = page._export_xlsx_to(str(out_file))

    assert result is False
    assert not out_file.exists()


def test_movements_report_xlsx_export_has_money_format_on_cout_unitaire(qtbot, login_as, tmp_path) -> None:
    import openpyxl

    stack, _ = login_as("Administrateur")
    supplier = stack.suppliers.create_supplier("F")
    article = _make_article(stack, stock_initial=Decimal("0"))
    entry = stack.entries.create_entry(supplier.id, date(2026, 1, 1), [EntreeLigneInput(article.id, Decimal("10"), Decimal("100"))])
    stack.entries.validate_entry(entry.id)

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Mouvements")

    out_file = tmp_path / "export.xlsx"
    page._export_xlsx_to(str(out_file))

    workbook = openpyxl.load_workbook(out_file)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    cout_col = headers.index("Coût unitaire") + 1
    assert sheet.cell(row=2, column=cout_col).data_type == "n"
    assert float(sheet.cell(row=2, column=cout_col).value) == 100


def test_reports_never_modify_stock(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    article = _make_article(stack, stock_initial=Decimal("42"))

    page = _build_page(stack)
    qtbot.addWidget(page)
    for report_name in ["État du stock", "Stock faible", "Ruptures", "Mouvements", "Entrées", "Sorties", "Ventes", "Inventaires", "Valorisation"]:
        page.report_combo.setCurrentText(report_name)

    assert stack.articles.get_article(article.id).stock_actuel == Decimal("42")


# -- Lot E-B.3 : OptionalDateEdit pour la période -----------------------------------------


def test_date_fields_are_optional_date_edit(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert isinstance(page.date_from_edit, OptionalDateEdit)
    assert isinstance(page.date_to_edit, OptionalDateEdit)


def test_dates_are_none_by_default(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    assert page.date_from_edit.date_or_none() is None
    assert page.date_to_edit.date_or_none() is None
    assert page._read_period() == (None, None)


def test_checking_date_from_makes_it_active(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.date_from_edit.set_date_or_none(date(2026, 1, 1))

    assert page.date_from_edit.checkbox.isChecked() is True
    assert page.date_from_edit.date_or_none() == date(2026, 1, 1)


def test_unchecking_date_from_returns_none(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)
    page.date_from_edit.set_date_or_none(date(2026, 1, 1))

    page.date_from_edit.checkbox.setChecked(False)

    assert page.date_from_edit.date_or_none() is None


def test_checking_date_to_makes_it_active(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.date_to_edit.set_date_or_none(date(2026, 1, 31))

    assert page.date_to_edit.checkbox.isChecked() is True
    assert page.date_to_edit.date_or_none() == date(2026, 1, 31)


def test_unchecking_date_to_returns_none(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)
    page.date_to_edit.set_date_or_none(date(2026, 1, 31))

    page.date_to_edit.checkbox.setChecked(False)

    assert page.date_to_edit.date_or_none() is None


def test_read_period_returns_selected_optional_dates(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.date_from_edit.set_date_or_none(date(2026, 1, 1))
    page.date_to_edit.set_date_or_none(date(2026, 1, 31))

    assert page._read_period() == (date(2026, 1, 1), date(2026, 1, 31))


def test_read_period_returns_none_for_unchecked_bound_only(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.date_from_edit.set_date_or_none(date(2026, 1, 1))
    # date_to reste décochée.

    assert page._read_period() == (date(2026, 1, 1), None)


def test_reset_filters_sets_both_dates_to_none(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)
    page.date_from_edit.set_date_or_none(date(2026, 1, 1))
    page.date_to_edit.set_date_or_none(date(2026, 1, 31))

    page._on_reset_filters_clicked()

    assert page.date_from_edit.date_or_none() is None
    assert page.date_to_edit.date_or_none() is None


def test_date_fields_hidden_for_report_without_period(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.report_combo.setCurrentText("État du stock")

    assert page.date_from_label.isHidden() is True
    assert page.date_from_edit.isHidden() is True
    assert page.date_to_label.isHidden() is True
    assert page.date_to_edit.isHidden() is True


def test_date_fields_visible_for_report_with_period(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    page.report_combo.setCurrentText("Mouvements")

    assert page.date_from_label.isHidden() is False
    assert page.date_from_edit.isHidden() is False
    assert page.date_to_label.isHidden() is False
    assert page.date_to_edit.isHidden() is False


def test_period_report_transmits_selected_dates_to_service(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    supplier = stack.suppliers.create_supplier("F")
    article = _make_article(stack, stock_initial=Decimal("0"))
    old_entry = stack.entries.create_entry(
        supplier.id, date(2020, 1, 1), [EntreeLigneInput(article.id, Decimal("5"), Decimal("100"))]
    )
    stack.entries.validate_entry(old_entry.id)
    recent_entry = stack.entries.create_entry(
        supplier.id, date(2026, 6, 15), [EntreeLigneInput(article.id, Decimal("5"), Decimal("100"))]
    )
    stack.entries.validate_entry(recent_entry.id)

    page = _build_page(stack)
    qtbot.addWidget(page)
    page.report_combo.setCurrentText("Entrées")
    assert page.table.rowCount() == 2  # aucune borne : tout l'historique

    page.date_from_edit.set_date_or_none(date(2026, 1, 1))
    page.date_to_edit.set_date_or_none(date(2026, 12, 31))
    page.refresh()

    assert page.table.rowCount() == 1
    assert page.table.item(0, 0).text() == recent_entry.numero


# -- correctif revue finale (Point 3) : voile non affiché pendant la construction --
#
# ReportsPage.__init__ déclenche un premier ``refresh()`` automatique
# (sélection du rapport par défaut) avant même que la page ne soit visible
# — potentiellement pendant que MainWindow construit toutes ses pages au
# démarrage (voir la revue de ce lot). Afficher le voile à ce moment précis
# n'apporte rien à l'utilisateur et ajoute un appel ``processEvents()``
# superflu pendant cette construction déjà sensible.


def test_overlay_is_not_shown_during_the_automatic_construction_time_refresh(qtbot, login_as, monkeypatch) -> None:
    stack, _ = login_as("Administrateur")
    calls = []
    monkeypatch.setattr(
        "app.views.pages.reports_page.LoadingOverlay.show_with_message",
        lambda self, message: calls.append(message),
    )

    page = _build_page(stack)
    qtbot.addWidget(page)

    assert calls == []


def test_overlay_is_shown_on_an_explicit_refresh_after_construction(qtbot, login_as) -> None:
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    calls = []
    page._loading_overlay.show_with_message = lambda message: calls.append(message)
    page._loading_overlay.hide_overlay = lambda: calls.append("hide")

    page.refresh()

    assert calls == ["Génération du rapport…", "hide"]


def test_overlay_is_shown_when_changing_the_report_type_after_construction(qtbot, login_as) -> None:
    """Le changement de type de rapport passe par le signal Qt
    ``currentIndexChanged`` (jamais un appel direct) : vérifie que
    ``show_loading`` reprend sa valeur par défaut (``True``) une fois la
    page construite, et non la valeur ``False`` utilisée une seule fois à
    la construction."""
    stack, _ = login_as("Administrateur")
    page = _build_page(stack)
    qtbot.addWidget(page)

    calls = []
    page._loading_overlay.show_with_message = lambda message: calls.append(message)

    page.report_combo.setCurrentText("Stock faible")

    assert calls == ["Génération du rapport…"]
